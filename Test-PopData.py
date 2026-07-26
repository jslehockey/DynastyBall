import openpyxl
import random
import os
import sqlite3
import pandas as pd
from Pfactory import PlayerFactory  # IMPORTING YOUR FACTORY

def populate_dummy_data(file_path):
    print("--- Step 1: Checking Excel Data ---")
    if not os.path.exists(file_path):
        print(f"Error: {file_path} not found. Run the schema builder first.")
        return False

    wb = openpyxl.load_workbook(file_path)
    
    # FIX 1: Clear existing data rows to force a fresh generation every time
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)

    # Re-populate Worlds
    wb["Worlds"].append(["Alpha", "Alpha Universe", "Yes"])

    # Re-populate Teams
    teams_data = [
        ("Alpha-0001", "Washington", "DiamondBucs", "Jim", "Active"),
        ("Alpha-0002", "Boston", "Minutemen", "Jake", "Active"),
        ("Alpha-0003", "Miami", "Stormers", "Rafael", "Active"),
        ("Alpha-0004", "Vaduz", "Mountaineers", "ETR", "Active")
    ]
    
    team_ids = []
    ws_teams = wb["Teams"]
    for t in teams_data:
        team_ids.append(t[0])
        ws_teams.append([t[0], "Alpha", t[1], t[2], t[3], t[4], f"{t[3]}@test.com", "pw123", "Active", "#000000", "#FFFFFF", 50])

    # Re-populate Parks
    ws_parks = wb["Parks"]
    for idx, t_id in enumerate(team_ids):
        ws_parks.append([f"Park-{idx+1}", t_id, f"Stadium {idx+1}", 5, 200, 150, 1000000] + [random.randint(12, 400) for _ in range(14)])

    # Re-populate Financials
    ws_fin = wb["Team_Financials"]
    for t_id in team_ids:
        ws_fin.append([t_id, 1, 15000000, 25, 4500000, 0, 2000000, 500000, 8000000])

    # Re-populate Players
    ws_base = wb["Players_Base"]
    ws_ratings = wb["Player_Ratings"]
    ws_contracts = wb["Contracts"]
    ws_hit_stats = wb["Stats_Hitting"]
    ws_pit_stats = wb["Stats_Pitching"]
    
    current_season = 1 
    factory = PlayerFactory(current_season)
    
    for t_id in team_ids:
        # GENERATE THE FULL ORGANIZATION USING YOUR FACTORY
        org = factory.generate_organization(team_name=t_id, is_expansion=True, league_tier=1)
        all_players = org["majors"] + org["minors"]
        
        for p in all_players:
            # 1. Base Data
            name_parts = p.name.split(' ', 1)
            fname = name_parts[0]
            lname = name_parts[1] if len(name_parts) > 1 else ""
            
            dev = p.attributes["development"]
            ws_base.append([p.player_id, fname, lname, dev["peak_age"], dev["last_peak_age"], dev["archetype"]])
            
            # 2. Ratings & Traits
            t_count = len(p.traits)
            t1 = p.traits[0] if t_count > 0 else ""
            t2 = p.traits[1] if t_count > 1 else ""
            t3 = p.traits[2] if t_count > 2 else ""

            recent_array = [random.randint(-4, 4) for _ in range(5)]
            recent_str = ", ".join(map(str, recent_array))
            psyche_mod = sum(recent_array)
            
            bat = p.attributes["batting"]
            run = p.attributes["baserunning"]
            df  = p.attributes["defense"]
            pit = p.attributes["pitching"]
            
            # Extracting specific stats depending on if they are a pitcher or hitter
            is_pitcher = p.attributes["Primary Pos"] == "P"
            stam_max = pit["stamina"] if is_pitcher else bat["stamina"]
            
            assigned_pos = p.assigned_pos if p.batting_order <= 9 or is_pitcher else "BENCH"

            # FIX 2: Safely map parent pitching stats. Falls back to sub-stats if parent keys don't exist.
            pit_velo = pit.get("velocity", pit.get("arm_speed", 0))
            pit_ctrl = pit.get("control", pit.get("accuracy", 0))
            pit_mov  = pit.get("movement", pit.get("spin_rate", 0))

            ratings_row = [
                p.player_id, current_season, t_id, p.league_level, dev["age"], p.attributes["Primary Pos"], "Active", 
                t1, t2, t3, t_count, recent_str, psyche_mod,
                assigned_pos, p.assigned_role, p.batting_order,
                # Hitting Stats
                bat["timing"], bat["barreling"], bat["strength"], bat["bat_speed"], bat["elevation"], 
                bat["eye"], bat["restraint"], run["sprint_speed"], run["instincts"], 
                # Defense
                df["def.range"], df["def.reaction"], df["def.glove"], df["def.ArmStr"], df["def.ArmAcc"],
                # Pitching/Stamina
                stam_max, stam_max, 
                pit_velo, pit["arm_speed"], pit["deception"], 
                pit_ctrl, pit["accuracy"], pit["command"], 
                pit_mov, pit["spin_rate"], pit["bite"]
            ]

            ws_ratings.append(ratings_row)
            
            # 3. Dummy Contract & Stats mapping (unchanged logic)
            ws_contracts.append([f"C-{p.player_id}", p.player_id, t_id, 1, 4, random.randint(50000, 2000000), "Active"])
            
            if is_pitcher:
                ip = random.randint(50, 200)
                er = int(ip * (random.uniform(2.5, 5.5) / 9))
                k = int(ip * random.uniform(0.7, 1.2))
                ws_pit_stats.append([p.player_id, current_season, "Regular Season", 30, random.randint(5, 20), random.randint(5, 15), 0, 0, 0, ip, int(ip*0.9), er+random.randint(0,5), er, random.randint(5, 25), random.randint(15, 60), random.randint(0, 5), k, ip*15, 1, 0, round((er*9)/ip, 2), 1.25, 3.85])
            else:
                ab = random.randint(300, 600)
                h = int(ab * random.uniform(0.220, 0.330))
                hr = random.randint(5, 40)
                ws_hit_stats.append([p.player_id, current_season, "Regular Season", 150, ab+50, ab, random.randint(40, 100), h, h-(hr+25), 20, 5, hr, random.randint(30, 100), random.randint(20, 80), random.randint(0, 10), random.randint(50, 150), random.randint(0, 30), random.randint(0, 10), round(h/ab, 3), 0.350, 0.450, 0.800])

    wb.save(file_path)
    print("[+] Dummy data generated and saved to Excel.")
    return True

def convert_excel_to_sqlite(excel_file, db_file):
    print("\n--- Step 2: Building SQLite Database ---")
    conn = sqlite3.connect(db_file)
    excel_data = pd.ExcelFile(excel_file)
    
    for sheet_name in excel_data.sheet_names:
        df = pd.read_excel(excel_file, sheet_name=sheet_name)
        if len(df.columns) == 0 or "Unnamed" in str(df.columns[0]):
            continue
        print(f"[*] Importing tab: {sheet_name}...")
        df.to_sql(sheet_name, conn, if_exists='replace', index=False)
        
    conn.close()
    print(f"\nSuccess! {db_file} is fully built and ready to query.")

if __name__ == "__main__":
    excel_filename = "DynastyBaseballSim.xlsx"
    db_filename = "diamondbucs_test.db"
    if populate_dummy_data(excel_filename):
        convert_excel_to_sqlite(excel_filename, db_filename)