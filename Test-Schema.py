import openpyxl
import os

def initialize_excel_schema(file_path):
    # 1. Verify the file exists before trying to open it
    if not os.path.exists(file_path):
        print(f"Error: Could not find {file_path} in the current directory.")
        return

    # 2. Load the existing workbook
    wb = openpyxl.load_workbook(file_path)
    
    # 3. Define our database structure (Tabs and their Columns)
    # You can easily add to this dictionary as the game expands
    schemas = {
        "Worlds": ["world_id", "world_name", "is_active"],
        
        "Teams": ["team_id", "world_id", "location", "nickname", "manager", "status", "user_email", "user_pw", "sub_status", "color_primary", "color_secondary", "prestige"],
        
        "Parks": ["park_id", "team_id", "name", "seasons_played", "record_w", "record_l", "population", "dim_lf_line", "dim_lf_height", "dim_dead_lf", "dim_dead_lf_height", "dim_lc_gap", "dim_lc_height", "dim_dead_center", "dim_dead_center_height", "dim_rc_gap", "dim_rc_height", "dim_dead_rf", "dim_dead_rf_height", "dim_rf_line", "dim_rf_height"],
        
        "Players_Base": ["player_id", "first_name", "last_name", "peak_age", "degrade_age", "archetype"],
        
       "Player_Ratings": [
            "player_id", "season", "team_id", "age", "position", "role", 
            "trait_1", "trait_2", "trait_3", "trait_count", 
            "recent_form", "psyche_mod",
            "con_timing", "con_barrel", "pow_str", "pow_batspd", "pow_elev", 
            "disc_eye", "disc_restr", "spd_sprint", "spd_inst", 
            "def_range", "def_react", "def_glove", "def_armstr", "def_armacc"
        ],
        
        "Stats_Hitting": ["player_id", "season", "competition", "g", "pa", "ab", "r", "h", "1b", "2b", "3b", "hr", "rbi", "bb", "hbp", "k", "sb", "cs", "avg", "obp", "slg", "ops"],

        "Stats_Pitching": ["player_id", "season", "competition", "g", "w", "l", "sv", "hld", "bs", "ip", "h", "r", "er", "hr", "bb", "hbp", "k", "pitches", "cg", "sho", "era", "whip", "fip"],
        
        "Stats_Fielding": ["player_id", "season", "competition", "g", "po", "a", "e", "tc", "fpct"],
        
        "Season_Results": ["game_id", "season", "day", "competition", "away_team_id", "away_score", "home_team_id", "home_score"],
        
        "Game_Events": ["event_id", "game_id", "inning", "half", "player_id", "event_type", "description"],

        "Team_Financials": ["team_id", "season", "balance", "players_on_contract", "cost_players", "income_fa_cup", "income_media", "income_adwatch", "income_attendance"],
        
        "Contracts": ["contract_id", "player_id", "team_id", "year_start", "year_end", "cost_per_season", "status"]
    }

    # 4. Loop through our desired schema and build the Excel file
    for tab_name, columns in schemas.items():
        if tab_name in wb.sheetnames:
            print(f"[*] Tab '{tab_name}' already exists. Skipping creation.")
            # If you later want to write dummy data to existing tabs, that code goes here.
        else:
            print(f"[+] Creating tab '{tab_name}' and adding schema headers...")
            ws = wb.create_sheet(title=tab_name)
            ws.append(columns) # Adds the list of columns as the first row

    # 5. Save the workbook with the new tabs
    wb.save(file_path)
    print("\nSuccess! Schema update complete. Open DynastyBaseballSim.xlsx to check it out.")

if __name__ == "__main__":
    # Ensure this matches your exact file name
    initialize_excel_schema("DynastyBaseballSim.xlsx")